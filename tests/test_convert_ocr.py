"""Tests for convert (PDF<->office), ocr (searchable scans), and the dependency doctor.

These features rely on system tools (LibreOffice, Tesseract) and optional pip extras.
Each test skips cleanly when its tools are missing, so the suite passes anywhere; where the
tools ARE installed (as in local dev), they run for real.
"""
import importlib.util
import io
import shutil

import pytest
from reportlab.pdfgen import canvas

import pdfblah as pb
from pdfblah import deps


def _has(mod):
    return importlib.util.find_spec(mod) is not None


HAS_LIBREOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))
HAS_TESSERACT = bool(shutil.which("tesseract"))
HAS_PDF2DOCX = _has("pdf2docx")
HAS_OCRMYPDF = _has("ocrmypdf")
HAS_DOCX = _has("docx")


def _make_pdf(path, text="Quarterly Report heading"):
    c = canvas.Canvas(str(path), pagesize=(400, 300))
    c.setFont("Helvetica", 16)
    c.drawString(40, 220, text)
    c.drawString(40, 180, "some body text here")
    c.save()
    return str(path)


# ---------- deps / doctor (always runnable) ----------
def test_deps_check_structure():
    st = deps.check()
    assert "tools" in st and "libs" in st
    for name in ("tesseract", "ghostscript", "libreoffice"):
        assert name in st["tools"]
        assert set(["present", "for"]).issubset(st["tools"][name])
    for mod in ("ocrmypdf", "pdf2docx"):
        assert mod in st["libs"]
    # summary lines render without raising
    assert isinstance(deps.summary_lines(), list) and deps.summary_lines()


def test_deps_install_hint():
    hint = deps.install_hint("tesseract")
    assert "tesseract" in hint and ("apt" in hint or "brew" in hint)


# ---------- convert ----------
@pytest.mark.skipif(not (HAS_LIBREOFFICE and HAS_DOCX), reason="needs LibreOffice + python-docx")
def test_office_to_pdf(tmp_path):
    from docx import Document
    src = tmp_path / "in.docx"
    d = Document(); d.add_heading("Hello Convert", 0); d.add_paragraph("body line"); d.save(str(src))
    out = tmp_path / "out.pdf"
    r = pb.convert(str(src), str(out))
    assert r["ok"], r
    import pdfplumber
    txt = pdfplumber.open(str(out)).pages[0].extract_text() or ""
    assert "Hello Convert" in txt


@pytest.mark.skipif(not HAS_PDF2DOCX, reason="needs pdf2docx")
def test_pdf_to_word(tmp_path):
    src = _make_pdf(tmp_path / "in.pdf", text="Convert Me To Word")
    out = tmp_path / "out.docx"
    r = pb.convert(src, str(out))
    assert r["ok"] and r["format"] == "docx", r
    assert out.exists() and out.stat().st_size > 0


def test_convert_unsupported_pair_is_refused(tmp_path):
    # PDF -> PNG is explicitly redirected to `render`, never a silent wrong guess
    r = pb.convert("x.pdf", "y.png")
    assert not r["ok"] and "render" in r["error"]


# ---------- ocr ----------
@pytest.mark.skipif(not (HAS_TESSERACT and HAS_OCRMYPDF), reason="needs Tesseract + ocrmypdf")
def test_ocr_makes_scan_searchable(tmp_path):
    import img2pdf
    import pdfplumber
    from PIL import Image, ImageDraw, ImageFont
    # build an image-only "scan" (no text layer)
    im = Image.new("RGB", (1000, 240), "white")
    d = ImageDraw.Draw(im)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 44)
    except Exception:
        font = ImageFont.load_default()
    d.text((40, 90), "OCR TARGET 4242", fill="black", font=font)
    png = tmp_path / "scan.png"; im.save(str(png))
    scan = tmp_path / "scan.pdf"
    scan.write_bytes(img2pdf.convert(str(png)))
    assert not (pdfplumber.open(str(scan)).pages[0].extract_text() or "").strip()
    out = tmp_path / "ocr.pdf"
    r = pb.ocr(str(scan), str(out), lang="eng")
    assert r["ok"], r
    text = pdfplumber.open(str(out)).pages[0].extract_text() or ""
    assert "OCR" in text and "4242" in text


def test_ocr_missing_lib_message(monkeypatch, tmp_path):
    # if ocrmypdf isn't importable, ocr() returns a clean message, not a traceback
    if HAS_OCRMYPDF:
        pytest.skip("ocrmypdf is installed here")
    r = pb.ocr(str(tmp_path / "x.pdf"), str(tmp_path / "y.pdf"))
    assert not r["ok"] and "ocr" in r["error"].lower()


# ---------- modular languages (always runnable: network and Tesseract are faked) ----------
@pytest.fixture
def tessdata_home(tmp_path, monkeypatch):
    ocr_mod = importlib.import_module("pdfblah.ocr")
    home = tmp_path / "tessdata"
    monkeypatch.setattr(ocr_mod, "TESSDATA_HOME", str(home))
    return home


def _fake_download(monkeypatch, payload=b"x" * 64, status=None):
    import io
    import urllib.error
    import urllib.request

    def fake_urlopen(url, timeout=0):
        if status:
            raise urllib.error.HTTPError(url, status, "nope", None, io.BytesIO())
        return io.BytesIO(payload)

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)


def test_get_language_downloads_one_file(tessdata_home, monkeypatch):
    _fake_download(monkeypatch)
    r = pb.get_language("deu")
    assert r["ok"] and r["quality"] == "fast", r
    assert (tessdata_home / "deu.traineddata").read_bytes() == b"x" * 64
    ocr_mod = importlib.import_module("pdfblah.ocr")
    assert ocr_mod._home_languages() == ["deu"]


def test_get_language_best_flag(tessdata_home, monkeypatch):
    import urllib.request
    seen = {}

    def spy(url, timeout=0):
        seen["url"] = url
        raise OSError("stop here")

    monkeypatch.setattr(urllib.request, "urlopen", spy)
    r = pb.get_language("fra", best=True)
    assert not r["ok"] and "tessdata_best" in seen["url"] and "fra.traineddata" in seen["url"]


def test_get_language_rejects_junk_codes(tessdata_home):
    for bad in ("", "de u", "../../etc/passwd", "a" * 40, "deu;rm"):
        r = pb.get_language(bad)
        assert not r["ok"] and r["code"] == "bad_lang", bad


def test_get_language_unknown_code_is_friendly(tessdata_home, monkeypatch):
    _fake_download(monkeypatch, status=404)
    r = pb.get_language("zzz")
    assert not r["ok"] and r["code"] == "unknown_lang"
    assert "tessdata_fast" in r["error"] and not (tessdata_home / "zzz.traineddata").exists()


def test_ensure_langs_prefers_system_packs(tessdata_home, monkeypatch):
    ocr_mod = importlib.import_module("pdfblah.ocr")
    monkeypatch.setattr(ocr_mod, "_system_tessdata", lambda: ("/sys/tessdata", ["eng", "deu"]))
    env, missing = ocr_mod._ensure_langs(["eng", "deu"])
    assert env == {} and missing == []


def test_ensure_langs_uses_downloaded_files(tessdata_home, monkeypatch):
    ocr_mod = importlib.import_module("pdfblah.ocr")
    tessdata_home.mkdir(parents=True)
    (tessdata_home / "deu.traineddata").write_bytes(b"d")
    monkeypatch.setattr(ocr_mod, "_system_tessdata", lambda: (None, []))
    env, missing = ocr_mod._ensure_langs(["deu"])
    assert env == {"TESSDATA_PREFIX": str(tessdata_home)} and missing == []


def test_ensure_langs_merges_system_into_home(tessdata_home, monkeypatch, tmp_path):
    # eng lives in the system dir, deu was downloaded: both must be visible in one dir
    ocr_mod = importlib.import_module("pdfblah.ocr")
    sysdir = tmp_path / "sys"
    sysdir.mkdir()
    (sysdir / "eng.traineddata").write_bytes(b"e")
    tessdata_home.mkdir(parents=True)
    (tessdata_home / "deu.traineddata").write_bytes(b"d")
    monkeypatch.setattr(ocr_mod, "_system_tessdata", lambda: (str(sysdir), ["eng"]))
    env, missing = ocr_mod._ensure_langs(["eng", "deu"])
    assert missing == [] and env == {"TESSDATA_PREFIX": str(tessdata_home)}
    assert (tessdata_home / "eng.traineddata").exists()


def test_ensure_langs_names_what_is_missing(tessdata_home, monkeypatch):
    ocr_mod = importlib.import_module("pdfblah.ocr")
    monkeypatch.setattr(ocr_mod, "_system_tessdata", lambda: (None, ["eng"]))
    env, missing = ocr_mod._ensure_langs(["eng", "fra"])
    assert env == {} and missing == ["fra"]


@pytest.mark.skipif(not HAS_OCRMYPDF, reason="needs ocrmypdf importable")
def test_ocr_missing_lang_says_get_lang(tessdata_home, monkeypatch, tmp_path):
    ocr_mod = importlib.import_module("pdfblah.ocr")
    monkeypatch.setattr("shutil.which", lambda n: "/usr/bin/tesseract")
    monkeypatch.setattr(ocr_mod, "_system_tessdata", lambda: (None, ["eng"]))
    r = pb.ocr(str(tmp_path / "x.pdf"), str(tmp_path / "y.pdf"), lang="eng+jpn")
    assert not r["ok"] and r["code"] == "no_lang"
    assert "--get-lang jpn" in r["error"]


def test_cli_langs_and_get_lang(tessdata_home, monkeypatch, capsys):
    from pdfblah import cli_tools
    monkeypatch.setattr(cli_tools, "ocr_languages", lambda: ["deu", "eng"])
    assert cli_tools.TOOL_HANDLERS["ocr"](["--langs"]) == 0
    assert capsys.readouterr().out.splitlines() == ["deu", "eng"]
    _fake_download(monkeypatch)
    assert cli_tools.TOOL_HANDLERS["ocr"](["--get-lang", "deu,fra"]) == 0
    out = capsys.readouterr().out
    assert "got deu" in out and "got fra" in out
    ocr_mod = importlib.import_module("pdfblah.ocr")
    assert ocr_mod._home_languages() == ["deu", "fra"]
